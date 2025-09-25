import os
import sys
import logging
from typing import List, Iterator
import tqdm
from docx.table import _Cell, Table
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
from docx import Document, ImagePart
from PIL import Image
from io import BytesIO
import numpy as np
from langchain_core.documents import Document
from langchain_community.document_loaders import UnstructuredFileLoader
from .base_loader import BaseLoader  # 导入基础接口

# 配置日志
logger = logging.getLogger(__name__)


class CustomizedOcrDocLoader(UnstructuredFileLoader, BaseLoader):
    
    def __init__(self, file_path: str):
        UnstructuredFileLoader.__init__(self, file_path=file_path)
        BaseLoader.__init__(self, file_path=file_path)

    def _is_paragraph_end(self, text):
        if text.strip()[-1] in ["。", "？", ".", "?"]:
            return True
        return False

    def _get_elements(self) -> List:
        try:
            from docx import Document
            doc = Document(self.file_path)

            full_text = []
            
            for para in doc.paragraphs:
                if para.text.strip():  
                    full_text.append(para.text.strip())

            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        full_text.append(" ".join(row_text))
            
            combined_text = "\n\n".join(full_text)
            
            # 如果文本太短，尝试OCR处理图片
            if len(combined_text) < 100:
                logger.warning(f"文本内容较少({len(combined_text)}字符)，尝试OCR提取图片文字...")
                try:
                    from rag.module.loader.ocr import get_rapid_ocr
                    rapid_ocr = get_rapid_ocr()
                    
                    # 提取文档中的图片
                    for rel in doc.part.rels.values():
                        if "image" in rel.target_ref:
                            try:
                                image_part = rel.target_part
                                if hasattr(image_part, '_blob'):
                                    image = Image.open(BytesIO(image_part._blob))
                                    rapid_ocr_result, _ = rapid_ocr(np.array(image))
                                    if rapid_ocr_result:
                                        ocr_text = " ".join([line[1] for line in rapid_ocr_result])
                                        if ocr_text.strip():
                                            full_text.append(f"[图片文字]: {ocr_text.strip()}")
                            except Exception as e:
                                logger.warning(f"OCR处理图片失败: {e}")
                                continue
                    
                    combined_text = "\n\n".join(full_text)
                except Exception as e:
                    print(f"OCR处理失败: {e}")
            
            # 使用unstructured处理文本
            if combined_text.strip():
                from unstructured.partition.text import partition_text
                return partition_text(text=combined_text, **self.unstructured_kwargs)
            else:
                logger.warning("文档内容为空！")
                return []
                
        except Exception as e:
            logger.error(f"DOCX文档加载失败: {e}")
            import traceback
            logger.error(traceback.format_exc())
            return []

    def load(self) -> List[Document]:
        """实现基础接口的load方法"""
        return list(self.lazy_load())

    def lazy_load(self) -> Iterator[Document]:
        """实现基础接口的lazy_load方法"""
        elements = self._get_elements()
        metadata = self._get_metadata()
        text = "\n\n".join([str(el) for el in elements])
        yield Document(page_content=text, metadata=metadata)

    @classmethod
    def supported_extensions(cls) -> List[str]:
        """返回支持的文件扩展名"""
        return [".docx"]


